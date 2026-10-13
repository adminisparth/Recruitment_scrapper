from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class ResumeParseError(Exception):
    pass


def parse_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))

        pages = []

        for page in reader.pages:
            text = page.extract_text() or ""

            if text.strip():
                pages.append(text)

        return "\n\n".join(pages).strip()

    except Exception as exc:
        raise ResumeParseError(
            f"Could not parse PDF: {exc}"
        ) from exc


def parse_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # Also extract text from tables.
        for table in document.tables:
            for row in table.rows:
                cells = []

                for cell in row.cells:
                    text = cell.text.strip()

                    if text:
                        cells.append(text)

                if cells:
                    paragraphs.append(" | ".join(cells))

        return "\n".join(paragraphs).strip()

    except Exception as exc:
        raise ResumeParseError(
            f"Could not parse DOCX: {exc}"
        ) from exc


def parse_resume(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError(
            f"Unsupported file type: {extension}. "
            f"Only PDF and DOCX are supported."
        )

    if not data:
        raise ResumeParseError("The uploaded file is empty.")

    if extension == ".pdf":
        text = parse_pdf(data)

    else:
        text = parse_docx(data)

    if len(text.strip()) < 100:
        raise ResumeParseError(
            f"Very little readable text was extracted from {filename}. "
            "If this is a scanned/image-only resume, OCR is required."
        )

    # Prevent enormous resumes from being sent to the LLM.
    return text[:20000]